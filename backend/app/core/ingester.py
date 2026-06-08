import os
import fitz  # PyMuPDF
from docx import Document as DocxDocument
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document
from app.config import get_settings
from app.core.vectorstore import get_vectorstore


def extract_text_from_pdf(filepath: str) -> list[Document]:
    docs = []
    pdf = fitz.open(filepath)
    filename = os.path.basename(filepath)
    for page_num in range(len(pdf)):
        page = pdf[page_num]
        text = page.get_text("text").strip()
        if text:
            docs.append(Document(
                page_content=text,
                metadata={
                    "source": filename,
                    "page": page_num + 1,
                    "total_pages": len(pdf),
                    "file_type": "pdf",
                }
            ))
    pdf.close()
    return docs


def extract_text_from_docx(filepath: str) -> list[Document]:
    docx = DocxDocument(filepath)
    filename = os.path.basename(filepath)
    full_text = []
    for para in docx.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())
    # Also extract tables
    for table in docx.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                full_text.append(row_text)
    combined = "\n".join(full_text)
    if combined.strip():
        return [Document(
            page_content=combined,
            metadata={"source": filename, "file_type": "docx", "page": 1}
        )]
    return []


def ingest_file(filepath: str, collection_name: str = "default") -> int:
    settings = get_settings()
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        raw_docs = extract_text_from_pdf(filepath)
    elif ext in (".docx", ".doc"):
        raw_docs = extract_text_from_docx(filepath)
    else:
        raise ValueError(f"Format tidak didukung: {ext}. Gunakan PDF atau DOCX.")

    if not raw_docs:
        raise ValueError("Tidak ada teks yang dapat diekstrak dari dokumen.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.max_chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=["\n\n", "\n", ".", " ", ""],
    )
    chunks = splitter.split_documents(raw_docs)

    # Tambahkan metadata chunk index
    for i, chunk in enumerate(chunks):
        chunk.metadata["chunk_index"] = i

    vectorstore = get_vectorstore(collection_name)
    vectorstore.add_documents(chunks)

    return len(chunks)
